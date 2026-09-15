"""
Packing / Delivery Sheet generator.
Takes per-unit plate counts for a session and produces the kitchen packing sheet.
Engine is reusable for the real app (this is the core of the Excel export module).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import math

# ---------------------------------------------------------------- CONFIG
# Per-plate conversion factors (verified against the 28 Apr 2026 sheet).
# 'source' tells which plate count drives the dish.
def make_dishes(count_vessels):
    return [
        {"name": "Idly",       "type": "count",  "per_plate": 4,   "source": "idly",       "vessels": count_vessels},
        {"name": "Sambar",     "type": "weight", "per_plate_g": 60, "source": "idly+wada",  "vessel_kg": 15},
        {"name": "FG Chatney", "type": "weight", "per_plate_g": 80, "source": "idly+wada",  "vessel_kg": 15},
        {"name": "Briyani",    "type": "weight", "per_plate_g": 300,"source": "biryani",    "vessel_kg": 20},
        {"name": "KURMA",      "type": "weight", "per_plate_g": 120,"source": "biryani",    "vessel_kg": 15},
        {"name": "Wada",       "type": "count",  "per_plate": 4,   "source": "wada",       "vessels": count_vessels},
    ]

MORNING_VESSELS = [200, 160, 120, 80, 40]
EVENING_VESSELS = [200, 160, 120, 100, 40]

# ---------------------------------------------------------------- ENGINE
def plates_for(source, row):
    if source == "idly":       return row["idly"]
    if source == "biryani":    return row["biryani"]
    if source == "wada":       return row["wada"]
    if source == "idly+wada":  return row["idly"] + row["wada"]
    return 0

def fill_vessels(qty, vessels):
    """Greedy: as many of the largest vessel as possible, then round the
    remainder UP to the smallest vessel that covers it. Returns {size: count}."""
    out = {v: 0 for v in vessels}
    if qty <= 0:
        return out
    vessels = sorted(vessels, reverse=True)
    largest = vessels[0]
    n = qty // largest
    out[largest] += n
    rem = qty - n * largest
    if rem > 0:
        covering = [v for v in vessels if v >= rem]
        chosen = min(covering) if covering else largest
        out[chosen] += 1
    return out

def compute_dish(dish, row):
    plates = plates_for(dish["source"], row)
    if dish["type"] == "count":
        qty = dish["per_plate"] * plates
        return {"qty": qty, "vessels": fill_vessels(qty, dish["vessels"])}
    else:
        qty_kg = round(dish["per_plate_g"] * plates / 1000, 1)
        vc = round(qty_kg / dish["vessel_kg"], 1) if dish["vessel_kg"] else 0
        return {"qty": qty_kg, "vessel_count": vc}

# ---------------------------------------------------------------- STYLES
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
GRP_FILL = PatternFill("solid", fgColor="FCE4D6")
TOT_FILL = PatternFill("solid", fgColor="FFF2CC")
RED = Font(name="Arial", size=9, bold=True, color="C00000")
BOLD = Font(name="Arial", size=9, bold=True)
NORM = Font(name="Arial", size=9)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_cell(c, font=NORM, fill=None, align=CTR, border=True):
    c.font = font
    c.alignment = align
    if fill: c.fill = fill
    if border: c.border = BORDER

# ---------------------------------------------------------------- RENDER
def render_block(ws, start_row, title2, date_str, rows, dishes, count_vessels):
    r = start_row
    # column plan ----------------------------------------------------
    # A SNo | B Route | C Idly D Biryani E Wada (indent) | F Total | dish blocks
    col = 7  # first dish col (G)
    dish_cols = []
    for d in dishes:
        width = 1 + (len(count_vessels) if d["type"] == "count" else 1)
        dish_cols.append((d, col, col + width - 1)); col += width
    last_col = col - 1

    # title rows
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last_col)
    c = ws.cell(r, 3, "Touch Stone Foundation - Vishakhapatnam"); style_cell(c, BOLD); r += 1
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last_col)
    c = ws.cell(r, 3, title2); style_cell(c, BOLD); r += 1
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last_col)
    c = ws.cell(r, 3, date_str); style_cell(c, BOLD); r += 1

    grp_row, sub_row = r, r + 1
    # left group headers (merged across 2 header rows)
    for cc, txt in [(1, "S.No"), (2, "Route"), (6, "Total")]:
        ws.merge_cells(start_row=grp_row, start_column=cc, end_row=sub_row, end_column=cc)
        style_cell(ws.cell(grp_row, cc, txt), BOLD, HDR_FILL)
    ws.merge_cells(start_row=grp_row, start_column=3, end_row=grp_row, end_column=5)
    style_cell(ws.cell(grp_row, 3, "Indent"), BOLD, HDR_FILL)
    for cc, txt in [(3, "Idly"), (4, "Biryani"), (5, "Wada")]:
        style_cell(ws.cell(sub_row, cc, txt), BOLD, HDR_FILL)
    # dish group headers
    for d, c0, c1 in dish_cols:
        ws.merge_cells(start_row=grp_row, start_column=c0, end_row=grp_row, end_column=c1)
        style_cell(ws.cell(grp_row, c0, d["name"]), BOLD, GRP_FILL)
        style_cell(ws.cell(sub_row, c0, "Qty"), BOLD, HDR_FILL)
        if d["type"] == "count":
            for i, v in enumerate(d["vessels"]):
                style_cell(ws.cell(sub_row, c0 + 1 + i, v), RED, HDR_FILL)
        else:
            style_cell(ws.cell(sub_row, c0 + 1, d["vessel_kg"]), RED, HDR_FILL)
    r = sub_row + 1

    # data rows
    totals = {"idly": 0, "biryani": 0, "wada": 0, "total": 0}
    dish_tot = {d["name"]: {"qty": 0, "vessels": {v: 0 for v in count_vessels}, "vc": 0} for d in dishes}
    for i, row in enumerate(rows, 1):
        style_cell(ws.cell(r, 1, i))
        style_cell(ws.cell(r, 2, row["route"]), NORM, align=Alignment(horizontal="left", vertical="center"))
        tot = row["idly"] + row["biryani"] + row["wada"]
        for cc, key in [(3, "idly"), (4, "biryani"), (5, "wada")]:
            style_cell(ws.cell(r, cc, row[key] or None)); totals[key] += row[key]
        style_cell(ws.cell(r, 6, tot), BOLD, TOT_FILL); totals["total"] += tot
        for d, c0, c1 in dish_cols:
            res = compute_dish(d, row)
            style_cell(ws.cell(r, c0, res["qty"] or None), BOLD)
            dish_tot[d["name"]]["qty"] += res["qty"]
            if d["type"] == "count":
                for j, v in enumerate(d["vessels"]):
                    n = res["vessels"][v]
                    style_cell(ws.cell(r, c0 + 1 + j, n or None))
                    dish_tot[d["name"]]["vessels"][v] += n
            else:
                style_cell(ws.cell(r, c0 + 1, res["vessel_count"] or None))
                dish_tot[d["name"]]["vc"] += res["vessel_count"]
        r += 1

    # totals row
    style_cell(ws.cell(r, 2, "Total"), BOLD, TOT_FILL,
               Alignment(horizontal="left", vertical="center"))
    style_cell(ws.cell(r, 1, None), BOLD, TOT_FILL)
    for cc, key in [(3, "idly"), (4, "biryani"), (5, "wada"), (6, "total")]:
        style_cell(ws.cell(r, cc, totals[key]), BOLD, TOT_FILL)
    for d, c0, c1 in dish_cols:
        dt = dish_tot[d["name"]]
        style_cell(ws.cell(r, c0, round(dt["qty"], 1)), BOLD, TOT_FILL)
        if d["type"] == "count":
            for j, v in enumerate(d["vessels"]):
                style_cell(ws.cell(r, c0 + 1 + j, dt["vessels"][v] or None), BOLD, TOT_FILL)
        else:
            style_cell(ws.cell(r, c0 + 1, round(dt["vc"], 1)), BOLD, TOT_FILL)
    return r + 2, last_col

def generate(path, date_str, morning_rows, evening_rows):
    wb = Workbook(); ws = wb.active; ws.title = "Delivery Sheet"
    r = 1
    r, lc = render_block(ws, r, "Brandix Morning Delivery Sheet", date_str,
                         morning_rows, make_dishes(MORNING_VESSELS), MORNING_VESSELS)
    r, lc = render_block(ws, r, "Brandix Evening Delivery Sheet", date_str,
                         evening_rows, make_dishes(EVENING_VESSELS), EVENING_VESSELS)
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 11
    for i in range(3, lc + 1):
        ws.column_dimensions[get_column_letter(i)].width = 6.5
    ws.freeze_panes = "A1"
    wb.save(path)
    return path

# ---------------------------------------------------------------- SAMPLE DATA (28 Apr 2026)
if __name__ == "__main__":
    morning = [
        {"route": "Unit - 1", "idly": 150, "biryani": 450, "wada": 150},
        {"route": "Unit - 2", "idly": 250, "biryani": 650, "wada": 200},
        {"route": "Unit - 3", "idly": 250, "biryani": 650, "wada": 180},
        {"route": "Unit - 4", "idly": 150, "biryani": 450, "wada": 150},
        {"route": "Vishaka",  "idly": 10,  "biryani": 16,  "wada": 20},
        {"route": "Others",   "idly": 10,  "biryani": 13,  "wada": 30},
    ]
    evening = [
        {"route": "Unit - 1", "idly": 80,  "biryani": 420, "wada": 50},
        {"route": "Unit - 2", "idly": 125, "biryani": 676, "wada": 125},
        {"route": "Unit - 3", "idly": 200, "biryani": 620, "wada": 130},
        {"route": "Unit - 4", "idly": 200, "biryani": 350, "wada": 100},
        {"route": "Vishaka",  "idly": 0,   "biryani": 34,  "wada": 0},
    ]
    out = generate("/sessions/cool-wizardly-dijkstra/mnt/outputs/packing_sheet_28Apr2026.xlsx",
                   "Tuesday, 28 April 2026", morning, evening)
    print("written:", out)

    # ---- verification dump vs original sheet
    print("\n== MORNING verification ==")
    for row in morning:
        dishes = make_dishes(MORNING_VESSELS)
        line = [row["route"]]
        for d in dishes:
            res = compute_dish(d, row)
            if d["type"] == "count":
                vv = ",".join(f"{k}:{n}" for k, n in res["vessels"].items() if n)
                line.append(f"{d['name']}={res['qty']} [{vv}]")
            else:
                line.append(f"{d['name']}={res['qty']}kg(/{d['vessel_kg']}={res['vessel_count']})")
        print(" | ".join(line))
